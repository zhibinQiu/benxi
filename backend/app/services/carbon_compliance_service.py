"""控排企业履约策略：CRUD、配置、策略运行门面。"""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timezone
from typing import Any

from sqlalchemy.orm import Session

from app.models.carbon_compliance import (
    SETTINGS_SINGLETON_ID,
    CarbonAlert,
    CarbonCeaHolding,
    CarbonCeaTrade,
    CarbonCcerHolding,
    CarbonCcerTrade,
    CarbonEmissionForecast,
    CarbonEmissionYear,
    CarbonEnterprise,
    CarbonGreenCert,
    CarbonGreenPower,
    CarbonMarketCeaMonthly,
    CarbonMarketCcerMonthly,
    CarbonMarketEnergyMonthly,
    CarbonStrategyRun,
    PlatformCarbonStrategySettings,
)
from app.services.carbon_compliance.accounting import AccountingInput, compute_accounting
from app.services.carbon_compliance.alerts import build_alerts
from app.services.carbon_compliance.carry_forward import compute_carry_forward
from app.services.carbon_compliance.compliance import eligible_ccer_qty
from app.services.carbon_compliance.defaults import (
    INDUSTRIES,
    INDUSTRY_LABELS,
    RISK_PROFILES,
    deep_merge,
    default_settings,
)
from app.services.carbon_compliance.market_cycle import judge_market_cycle
from app.services.carbon_compliance.market_sync import (
    CNEEEX_DAILY_URL,
    DailyQuote,
    fetch_structured_quotes,
    period_tag_for_month,
)
from app.services.carbon_compliance.report_export import strategy_run_to_markdown
from app.services.carbon_compliance.strategy_engine import build_strategy_payload

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore
    load_workbook = None  # type: ignore


def get_settings(db: Session) -> dict:
    row = db.get(PlatformCarbonStrategySettings, SETTINGS_SINGLETON_ID)
    if not row:
        row = PlatformCarbonStrategySettings(
            id=SETTINGS_SINGLETON_ID, payload=default_settings()
        )
        db.add(row)
        db.commit()
        db.refresh(row)
    return deep_merge(default_settings(), row.payload or {})


def update_settings(db: Session, payload: dict) -> dict:
    row = db.get(PlatformCarbonStrategySettings, SETTINGS_SINGLETON_ID)
    if not row:
        row = PlatformCarbonStrategySettings(id=SETTINGS_SINGLETON_ID, payload={})
        db.add(row)
    merged = deep_merge(default_settings(), deep_merge(row.payload or {}, payload or {}))
    row.payload = merged
    db.commit()
    db.refresh(row)
    return deep_merge(default_settings(), row.payload or {})


def list_enterprises(db: Session, user_id: uuid.UUID) -> list[CarbonEnterprise]:
    return (
        db.query(CarbonEnterprise)
        .filter(CarbonEnterprise.user_id == user_id)
        .order_by(CarbonEnterprise.updated_at.desc())
        .all()
    )


def get_enterprise(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID
) -> CarbonEnterprise | None:
    ent = db.get(CarbonEnterprise, enterprise_id)
    if not ent or ent.user_id != user_id:
        return None
    return ent


def create_enterprise(db: Session, user_id: uuid.UUID, data: dict) -> CarbonEnterprise:
    industry = data.get("industry") or ""
    if industry not in INDUSTRIES:
        raise ValueError(f"industry must be one of {INDUSTRIES}")
    risk = data.get("risk_profile") or "balanced"
    if risk not in RISK_PROFILES:
        raise ValueError(f"risk_profile must be one of {RISK_PROFILES}")
    raw_name = str(data.get("name") or "").strip()
    if not raw_name:
        label = INDUSTRY_LABELS.get(industry, industry)
        uscc = str(data.get("uscc") or "").strip()
        raw_name = f"{label}履约主体" + (f"-{uscc[-6:]}" if uscc else "")
    ent = CarbonEnterprise(
        user_id=user_id,
        name=raw_name,
        uscc=str(data.get("uscc") or "").strip(),
        industry=industry,
        market_start_year=int(data["market_start_year"]),
        compliance_cycle=str(data.get("compliance_cycle") or "annual"),
        risk_profile=risk,
        annual_budget_cap=float(data.get("annual_budget_cap") or 0),
        single_trade_limit=float(data.get("single_trade_limit") or 0),
        enterprise_attrs=dict(data.get("enterprise_attrs") or {}),
    )
    db.add(ent)
    db.commit()
    db.refresh(ent)
    return ent


def update_enterprise(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, data: dict
) -> CarbonEnterprise:
    ent = get_enterprise(db, user_id, enterprise_id)
    if not ent:
        raise LookupError("enterprise not found")
    if "name" in data and data["name"] is not None:
        ent.name = str(data["name"]).strip()
    if "uscc" in data:
        ent.uscc = str(data.get("uscc") or "").strip()
    if "industry" in data and data["industry"] is not None:
        if data["industry"] not in INDUSTRIES:
            raise ValueError(f"industry must be one of {INDUSTRIES}")
        ent.industry = data["industry"]
    if "market_start_year" in data and data["market_start_year"] is not None:
        ent.market_start_year = int(data["market_start_year"])
    if "compliance_cycle" in data and data["compliance_cycle"] is not None:
        ent.compliance_cycle = str(data["compliance_cycle"])
    if "risk_profile" in data and data["risk_profile"] is not None:
        if data["risk_profile"] not in RISK_PROFILES:
            raise ValueError(f"risk_profile must be one of {RISK_PROFILES}")
        ent.risk_profile = data["risk_profile"]
    if "annual_budget_cap" in data and data["annual_budget_cap"] is not None:
        ent.annual_budget_cap = float(data["annual_budget_cap"])
    if "single_trade_limit" in data and data["single_trade_limit"] is not None:
        ent.single_trade_limit = float(data["single_trade_limit"])
    if "enterprise_attrs" in data and data["enterprise_attrs"] is not None:
        ent.enterprise_attrs = dict(data["enterprise_attrs"])
    db.commit()
    db.refresh(ent)
    return ent


def delete_enterprise(db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID) -> None:
    ent = get_enterprise(db, user_id, enterprise_id)
    if not ent:
        raise LookupError("enterprise not found")
    db.delete(ent)
    db.commit()


def _upsert_emission_year(db: Session, enterprise_id: uuid.UUID, data: dict) -> CarbonEmissionYear:
    year = int(data["year"])
    row = (
        db.query(CarbonEmissionYear)
        .filter(
            CarbonEmissionYear.enterprise_id == enterprise_id,
            CarbonEmissionYear.year == year,
        )
        .first()
    )
    if not row:
        row = CarbonEmissionYear(enterprise_id=enterprise_id, year=year)
        db.add(row)
    for field in (
        "verified_total",
        "scope1_combustion",
        "scope1_process",
        "scope2_power",
        "purchased_mwh",
        "historical_gap",
        "ccer_used",
    ):
        if field in data:
            setattr(row, field, data[field])
    if "monthly_detail" in data:
        row.monthly_detail = dict(data["monthly_detail"] or {})
    db.commit()
    db.refresh(row)
    return row


def list_emission_years(db: Session, enterprise_id: uuid.UUID) -> list[CarbonEmissionYear]:
    return (
        db.query(CarbonEmissionYear)
        .filter(CarbonEmissionYear.enterprise_id == enterprise_id)
        .order_by(CarbonEmissionYear.year.desc())
        .all()
    )


def upsert_emission_year(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, data: dict
) -> CarbonEmissionYear:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    return _upsert_emission_year(db, enterprise_id, data)


def delete_emission_year(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, year: int
) -> None:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    row = (
        db.query(CarbonEmissionYear)
        .filter(
            CarbonEmissionYear.enterprise_id == enterprise_id,
            CarbonEmissionYear.year == int(year),
        )
        .first()
    )
    if not row:
        raise LookupError("emission year not found")
    db.delete(row)
    db.commit()


def upsert_forecast(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, data: dict
) -> CarbonEmissionForecast:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    year = int(data["year"])
    row = (
        db.query(CarbonEmissionForecast)
        .filter(
            CarbonEmissionForecast.enterprise_id == enterprise_id,
            CarbonEmissionForecast.year == year,
        )
        .first()
    )
    if not row:
        row = CarbonEmissionForecast(enterprise_id=enterprise_id, year=year)
        db.add(row)
    row.forecast_total = float(data.get("forecast_total") or 0)
    row.capacity_plan = str(data.get("capacity_plan") or "")
    row.abatement_projects = list(data.get("abatement_projects") or [])
    row.production_plan = str(data.get("production_plan") or "")
    db.commit()
    db.refresh(row)
    return row


def list_forecasts(db: Session, enterprise_id: uuid.UUID) -> list[CarbonEmissionForecast]:
    return (
        db.query(CarbonEmissionForecast)
        .filter(CarbonEmissionForecast.enterprise_id == enterprise_id)
        .order_by(CarbonEmissionForecast.year.desc())
        .all()
    )


def upsert_cea_holding(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, data: dict
) -> CarbonCeaHolding:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    vy = int(data["vintage_year"])
    row = (
        db.query(CarbonCeaHolding)
        .filter(
            CarbonCeaHolding.enterprise_id == enterprise_id,
            CarbonCeaHolding.vintage_year == vy,
        )
        .first()
    )
    if not row:
        row = CarbonCeaHolding(enterprise_id=enterprise_id, vintage_year=vy)
        db.add(row)
    for field in (
        "free_quota",
        "carry_forward_qty",
        "net_sell_qty",
        "avg_cost",
        "estimated_free_quota",
        "sellable_cap",
    ):
        if field in data:
            setattr(row, field, data[field])
    # 兼容旧字段名
    if "carry_forward_qty" not in data and "expired_qty" in data:
        row.carry_forward_qty = float(data.get("expired_qty") or 0)
    db.commit()
    db.refresh(row)
    return row


def list_cea_holdings(db: Session, enterprise_id: uuid.UUID) -> list[CarbonCeaHolding]:
    return (
        db.query(CarbonCeaHolding)
        .filter(CarbonCeaHolding.enterprise_id == enterprise_id)
        .order_by(CarbonCeaHolding.vintage_year.desc())
        .all()
    )


def delete_cea_holding(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, vintage_year: int
) -> None:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    row = (
        db.query(CarbonCeaHolding)
        .filter(
            CarbonCeaHolding.enterprise_id == enterprise_id,
            CarbonCeaHolding.vintage_year == int(vintage_year),
        )
        .first()
    )
    if not row:
        raise LookupError("cea holding not found")
    db.delete(row)
    db.commit()


def get_cea_holding_for_year(
    db: Session, enterprise_id: uuid.UUID, vintage_year: int
) -> CarbonCeaHolding | None:
    return (
        db.query(CarbonCeaHolding)
        .filter(
            CarbonCeaHolding.enterprise_id == enterprise_id,
            CarbonCeaHolding.vintage_year == int(vintage_year),
        )
        .first()
    )

def add_cea_trade(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, data: dict
) -> CarbonCeaTrade:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    side = data.get("side")
    if side not in ("buy", "sell"):
        raise ValueError("side must be buy or sell")
    row = CarbonCeaTrade(
        enterprise_id=enterprise_id,
        side=side,
        qty=float(data["qty"]),
        price=float(data.get("price") or 0),
        note=str(data.get("note") or ""),
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def list_cea_trades(db: Session, enterprise_id: uuid.UUID) -> list[CarbonCeaTrade]:
    return (
        db.query(CarbonCeaTrade)
        .filter(CarbonCeaTrade.enterprise_id == enterprise_id)
        .order_by(CarbonCeaTrade.traded_at.desc())
        .all()
    )


def cea_net_sell_qty(db: Session, enterprise_id: uuid.UUID) -> float:
    """从交易流水汇总净卖出 = 卖出量 − 买入量（万吨）。"""
    sold = 0.0
    bought = 0.0
    for t in list_cea_trades(db, enterprise_id):
        q = float(t.qty or 0)
        if t.side == "sell":
            sold += q
        elif t.side == "buy":
            bought += q
    return sold - bought


def resolve_cea_net_sell(
    db: Session, enterprise_id: uuid.UUID, compliance_year: int
) -> float:
    """结转测算用净卖出：优先取履约年 CEA 台账「当前净卖出」，否则回退交易流水汇总。"""
    row = get_cea_holding_for_year(db, enterprise_id, compliance_year)
    if row is not None:
        return float(getattr(row, "net_sell_qty", 0) or 0)
    return cea_net_sell_qty(db, enterprise_id)


def upsert_ccer_holding(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, data: dict
) -> CarbonCcerHolding:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    expire_at = data.get("expire_at")
    if isinstance(expire_at, str) and expire_at:
        expire_at = date.fromisoformat(expire_at[:10])
    qty = max(0.0, float(data.get("qty") or 0))
    if data.get("eligible_qty") is not None:
        eligible = max(0.0, float(data["eligible_qty"]))
        if eligible <= 0 and qty > 0:
            eligible = qty
        elif qty > 0:
            eligible = min(eligible, qty)
    else:
        eligible = qty
    row = CarbonCcerHolding(
        enterprise_id=enterprise_id,
        project_type=str(data.get("project_type") or ""),
        issue_year=int(data["issue_year"]),
        expire_at=expire_at,
        qty=qty,
        cost=float(data.get("cost") or 0),
        eligible_qty=eligible,
        linked_green_cert=bool(data.get("linked_green_cert") or False),
    )
    if data.get("id"):
        existing = db.get(CarbonCcerHolding, uuid.UUID(str(data["id"])))
        if existing and existing.enterprise_id == enterprise_id:
            for k in (
                "project_type",
                "issue_year",
                "expire_at",
                "qty",
                "cost",
                "eligible_qty",
                "linked_green_cert",
            ):
                setattr(existing, k, getattr(row, k))
            db.commit()
            db.refresh(existing)
            return existing
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def list_ccer_holdings(db: Session, enterprise_id: uuid.UUID) -> list[CarbonCcerHolding]:
    return (
        db.query(CarbonCcerHolding)
        .filter(CarbonCcerHolding.enterprise_id == enterprise_id)
        .order_by(CarbonCcerHolding.issue_year.desc())
        .all()
    )


def delete_ccer_holding(
    db: Session,
    user_id: uuid.UUID,
    enterprise_id: uuid.UUID,
    holding_id: uuid.UUID,
) -> None:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    row = db.get(CarbonCcerHolding, holding_id)
    if not row or row.enterprise_id != enterprise_id:
        raise LookupError("ccer holding not found")
    db.delete(row)
    db.commit()


def add_ccer_trade(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, data: dict
) -> CarbonCcerTrade:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    side = data.get("side")
    if side not in ("buy", "sell"):
        raise ValueError("side must be buy or sell")
    row = CarbonCcerTrade(
        enterprise_id=enterprise_id,
        side=side,
        qty=float(data["qty"]),
        price=float(data.get("price") or 0),
        note=str(data.get("note") or ""),
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def upsert_green_power(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, data: dict
) -> CarbonGreenPower:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    year = int(data["year"])
    row = (
        db.query(CarbonGreenPower)
        .filter(
            CarbonGreenPower.enterprise_id == enterprise_id,
            CarbonGreenPower.year == year,
        )
        .first()
    )
    if not row:
        row = CarbonGreenPower(enterprise_id=enterprise_id, year=year)
        db.add(row)
    row.market_green_mwh = float(data.get("market_green_mwh") or 0)
    row.self_gen_mwh = float(data.get("self_gen_mwh") or 0)
    row.premium_per_mwh = float(data.get("premium_per_mwh") or 0)
    row.contract_ref = str(data.get("contract_ref") or "")
    db.commit()
    db.refresh(row)
    return row


def list_green_power(db: Session, enterprise_id: uuid.UUID) -> list[CarbonGreenPower]:
    return (
        db.query(CarbonGreenPower)
        .filter(CarbonGreenPower.enterprise_id == enterprise_id)
        .order_by(CarbonGreenPower.year.desc())
        .all()
    )


def upsert_green_cert(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, data: dict
) -> CarbonGreenCert:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    row = CarbonGreenCert(
        enterprise_id=enterprise_id,
        year=int(data["year"]),
        qty=float(data.get("qty") or 0),
        unit_price=float(data.get("unit_price") or 0),
        retired=bool(data.get("retired") or False),
        ren_weight_target=data.get("ren_weight_target"),
    )
    if data.get("id"):
        existing = db.get(CarbonGreenCert, uuid.UUID(str(data["id"])))
        if existing and existing.enterprise_id == enterprise_id:
            for k in ("year", "qty", "unit_price", "retired", "ren_weight_target"):
                setattr(existing, k, getattr(row, k))
            db.commit()
            db.refresh(existing)
            return existing
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def list_green_certs(db: Session, enterprise_id: uuid.UUID) -> list[CarbonGreenCert]:
    return (
        db.query(CarbonGreenCert)
        .filter(CarbonGreenCert.enterprise_id == enterprise_id)
        .order_by(CarbonGreenCert.year.desc())
        .all()
    )


def upsert_market_cea(db: Session, data: dict) -> CarbonMarketCeaMonthly:
    ym = str(data["year_month"])
    row = (
        db.query(CarbonMarketCeaMonthly)
        .filter(CarbonMarketCeaMonthly.year_month == ym)
        .first()
    )
    if not row:
        row = CarbonMarketCeaMonthly(year_month=ym, avg_price=float(data["avg_price"]))
        db.add(row)
    row.avg_price = float(data["avg_price"])
    row.high = data.get("high")
    row.low = data.get("low")
    row.period_tag = str(data.get("period_tag") or "")
    db.commit()
    db.refresh(row)
    return row


def list_market_cea(db: Session, limit: int = 60) -> list[CarbonMarketCeaMonthly]:
    return (
        db.query(CarbonMarketCeaMonthly)
        .order_by(CarbonMarketCeaMonthly.year_month.desc())
        .limit(limit)
        .all()
    )


def upsert_market_ccer(db: Session, data: dict) -> CarbonMarketCcerMonthly:
    ym = str(data["year_month"])
    ptype = str(data.get("project_type") or "general")
    row = (
        db.query(CarbonMarketCcerMonthly)
        .filter(
            CarbonMarketCcerMonthly.year_month == ym,
            CarbonMarketCcerMonthly.project_type == ptype,
        )
        .first()
    )
    if not row:
        row = CarbonMarketCcerMonthly(
            year_month=ym, project_type=ptype, avg_price=float(data["avg_price"])
        )
        db.add(row)
    row.avg_price = float(data["avg_price"])
    db.commit()
    db.refresh(row)
    return row


def list_market_ccer(db: Session, limit: int = 60) -> list[CarbonMarketCcerMonthly]:
    return (
        db.query(CarbonMarketCcerMonthly)
        .order_by(CarbonMarketCcerMonthly.year_month.desc())
        .limit(limit)
        .all()
    )


def upsert_market_energy(db: Session, data: dict) -> CarbonMarketEnergyMonthly:
    ym = str(data["year_month"])
    region = str(data.get("region") or "national")
    row = (
        db.query(CarbonMarketEnergyMonthly)
        .filter(
            CarbonMarketEnergyMonthly.year_month == ym,
            CarbonMarketEnergyMonthly.region == region,
        )
        .first()
    )
    if not row:
        row = CarbonMarketEnergyMonthly(year_month=ym, region=region)
        db.add(row)
    row.green_premium = data.get("green_premium")
    row.grec_price = data.get("grec_price")
    row.coal_price = data.get("coal_price")
    db.commit()
    db.refresh(row)
    return row


def list_market_energy(db: Session, limit: int = 60) -> list[CarbonMarketEnergyMonthly]:
    return (
        db.query(CarbonMarketEnergyMonthly)
        .order_by(CarbonMarketEnergyMonthly.year_month.desc())
        .limit(limit)
        .all()
    )


def _merge_cea_month(db: Session, quote: DailyQuote) -> CarbonMarketCeaMonthly:
    ym = quote.year_month
    price = quote.primary_price
    if price is None:
        raise ValueError("cea quote missing price")
    row = (
        db.query(CarbonMarketCeaMonthly)
        .filter(CarbonMarketCeaMonthly.year_month == ym)
        .first()
    )
    day_high = quote.high if quote.high is not None else price
    day_low = quote.low if quote.low is not None else price
    if not row:
        row = CarbonMarketCeaMonthly(
            year_month=ym,
            avg_price=float(price),
            high=float(day_high),
            low=float(day_low),
            period_tag=period_tag_for_month(quote.trade_date.month),
        )
        db.add(row)
    else:
        row.avg_price = float(price)
        prev_high = float(row.high) if row.high is not None else float(price)
        prev_low = float(row.low) if row.low is not None else float(price)
        row.high = max(prev_high, float(day_high))
        row.low = min(prev_low, float(day_low))
        row.period_tag = period_tag_for_month(quote.trade_date.month)
    return row


def _upsert_cea_month_agg(
    db: Session,
    *,
    year_month: str,
    avg_price: float,
    high: float,
    low: float,
    period_tag: str = "",
    use_last_close_as_avg: bool = False,
    last_close: float | None = None,
) -> CarbonMarketCeaMonthly:
    """按月全量覆盖写入（来自日线聚合）。"""
    row = (
        db.query(CarbonMarketCeaMonthly)
        .filter(CarbonMarketCeaMonthly.year_month == year_month)
        .first()
    )
    price = float(last_close if use_last_close_as_avg and last_close is not None else avg_price)
    if not row:
        row = CarbonMarketCeaMonthly(
            year_month=year_month,
            avg_price=price,
            high=float(high),
            low=float(low),
            period_tag=period_tag or period_tag_for_month(int(year_month.split("-")[1])),
        )
        db.add(row)
    else:
        row.avg_price = price
        row.high = float(high)
        row.low = float(low)
        row.period_tag = period_tag or row.period_tag
    return row


def _merge_ccer_month(db: Session, quote: DailyQuote) -> CarbonMarketCcerMonthly:
    ym = quote.year_month
    price = quote.primary_price
    if price is None:
        raise ValueError("ccer quote missing price")
    ptype = "general"
    row = (
        db.query(CarbonMarketCcerMonthly)
        .filter(
            CarbonMarketCcerMonthly.year_month == ym,
            CarbonMarketCcerMonthly.project_type == ptype,
        )
        .first()
    )
    if not row:
        row = CarbonMarketCcerMonthly(
            year_month=ym, project_type=ptype, avg_price=float(price)
        )
        db.add(row)
    else:
        row.avg_price = float(price)
    return row


async def sync_market_quotes(db: Session) -> dict[str, Any]:
    """外站拉取 CEA（环交所 K 线）/ CCER 并写入月度表。"""
    fetched = await fetch_structured_quotes()
    written: dict[str, Any] = {"cea": None, "ccer": None, "cea_months": 0}
    errors: list[str] = []

    monthly_rows = fetched.get("cea_monthly") or []
    if monthly_rows:
        try:
            # 历史月：均价用月内收盘均值；当月：用月末/最新收盘更贴近现价
            latest_ym = monthly_rows[-1]["year_month"] if monthly_rows else ""
            last_row = None
            for m in monthly_rows:
                is_latest = m["year_month"] == latest_ym
                last_row = _upsert_cea_month_agg(
                    db,
                    year_month=m["year_month"],
                    avg_price=float(m["avg_price"]),
                    high=float(m["high"]),
                    low=float(m["low"]),
                    period_tag=str(m.get("period_tag") or ""),
                    use_last_close_as_avg=is_latest,
                    last_close=float(m.get("last_close") or m["avg_price"]),
                )
            written["cea_months"] = len(monthly_rows)
            if last_row:
                written["cea"] = {
                    "year_month": last_row.year_month,
                    "avg_price": last_row.avg_price,
                    "high": last_row.high,
                    "low": last_row.low,
                    "trade_date": (fetched.get("cea") or {}).get("trade_date"),
                    "source": CNEEEX_DAILY_URL
                    if "cneeex.com" in str((fetched.get("sources_tried") or [""])[0])
                    else (fetched.get("cea") or {}).get("source"),
                    "daily_count": fetched.get("daily_count") or 0,
                }
        except Exception as exc:
            errors.append(f"cea_monthly: {exc}")
    elif fetched.get("cea"):
        try:
            raw = fetched["cea"]
            q = DailyQuote(
                trade_date=date.fromisoformat(raw["trade_date"]),
                open=raw.get("open"),
                high=raw.get("high"),
                low=raw.get("low"),
                close=raw.get("close"),
                avg_price=raw.get("avg_price"),
                source=raw.get("source") or "",
                instrument="cea",
            )
            row = _merge_cea_month(db, q)
            written["cea"] = {
                "year_month": row.year_month,
                "avg_price": row.avg_price,
                "high": row.high,
                "low": row.low,
                "trade_date": q.trade_date.isoformat(),
                "source": q.source,
            }
            written["cea_months"] = 1
        except Exception as exc:
            errors.append(f"cea: {exc}")

    if fetched.get("ccer"):
        try:
            raw = fetched["ccer"]
            q = DailyQuote(
                trade_date=date.fromisoformat(raw["trade_date"]),
                avg_price=raw.get("avg_price"),
                close=raw.get("close"),
                source=raw.get("source") or "",
                instrument="ccer",
            )
            row = _merge_ccer_month(db, q)
            written["ccer"] = {
                "year_month": row.year_month,
                "avg_price": row.avg_price,
                "project_type": row.project_type,
                "trade_date": q.trade_date.isoformat(),
                "source": q.source,
            }
        except Exception as exc:
            errors.append(f"ccer: {exc}")

    settings = get_settings(db)
    integrations = dict(settings.get("integrations") or {})
    integrations["last_sync_at"] = datetime.now(timezone.utc).astimezone().isoformat(
        timespec="seconds"
    )
    integrations["last_sync_ok"] = bool(written["cea"] or written["ccer"])
    integrations["last_sync_detail"] = {
        "sources_tried": fetched.get("sources_tried") or [],
        "written": written,
        "errors": errors,
        "daily_count": fetched.get("daily_count") or 0,
    }
    update_settings(db, {"integrations": integrations})
    db.commit()

    return {
        "ok": bool(written["cea"] or written["ccer"]),
        "queried_at": fetched.get("queried_at"),
        "sources_tried": fetched.get("sources_tried") or [],
        "written": written,
        "errors": errors,
        "raw": {
            "cea": fetched.get("cea"),
            "ccer": fetched.get("ccer"),
            "intraday": fetched.get("intraday"),
            "cea_months": len(monthly_rows),
        },
    }


def should_auto_sync_market(settings: dict | None = None) -> bool:
    """根据 market_sync_period 与 last_sync_at 判断是否应自动同步。"""
    cfg = (settings or {}).get("integrations") or {}
    if cfg.get("market_sync_enabled") is False:
        return False
    period = str(cfg.get("market_sync_period") or "day").lower()
    last = cfg.get("last_sync_at")
    if not last:
        return True
    try:
        last_dt = datetime.fromisoformat(str(last).replace("Z", "+00:00"))
    except Exception:
        return True
    now = datetime.now(timezone.utc)
    if last_dt.tzinfo is None:
        last_dt = last_dt.replace(tzinfo=timezone.utc)
    age_h = (now - last_dt.astimezone(timezone.utc)).total_seconds() / 3600.0
    if period in ("hour", "hourly"):
        return age_h >= 1.0
    if period in ("day", "daily"):
        return age_h >= 12.0
    return age_h >= 24.0


def _holding_to_dict(h: CarbonCcerHolding) -> dict:
    return {
        "eligible_qty": h.eligible_qty,
        "qty": h.qty,
        "expire_at": h.expire_at.isoformat() if h.expire_at else None,
        "linked_green_cert": h.linked_green_cert,
    }


def run_strategy(
    db: Session,
    user_id: uuid.UUID,
    enterprise_id: uuid.UUID,
    compliance_year: int,
    *,
    notify: bool = True,
) -> CarbonStrategyRun:
    ent = get_enterprise(db, user_id, enterprise_id)
    if not ent:
        raise LookupError("enterprise not found")
    settings = get_settings(db)
    industry_params = (settings.get("industry_params") or {}).get(ent.industry) or {}
    cost = settings.get("cost") or {}
    compliance_cfg = settings.get("compliance") or {}

    emission = (
        db.query(CarbonEmissionYear)
        .filter(
            CarbonEmissionYear.enterprise_id == enterprise_id,
            CarbonEmissionYear.year == compliance_year,
        )
        .first()
    )
    forecast = (
        db.query(CarbonEmissionForecast)
        .filter(
            CarbonEmissionForecast.enterprise_id == enterprise_id,
            CarbonEmissionForecast.year == compliance_year,
        )
        .first()
    )
    cea_rows = list_cea_holdings(db, enterprise_id)
    free_cea = 0.0
    sellable = 0.0
    allocated_free = 0.0  # 当年新分配（不含结转）
    carry_in = 0.0
    for c in cea_rows:
        if c.vintage_year == compliance_year or c.vintage_year >= ent.market_start_year:
            allocated = max(0.0, float(c.free_quota or 0))
            carried = max(0.0, float(getattr(c, "carry_forward_qty", 0) or 0))
            usable = allocated + carried
            free_cea += usable
            if c.vintage_year == compliance_year:
                allocated_free += allocated
                carry_in += carried
            if c.estimated_free_quota and c.vintage_year == compliance_year:
                free_cea = max(free_cea, float(c.estimated_free_quota) + carried)
            if c.sellable_cap is not None:
                sellable += float(c.sellable_cap)
            else:
                sellable += usable

    ccer_rows = list_ccer_holdings(db, enterprise_id)
    own_ccer = eligible_ccer_qty([_holding_to_dict(h) for h in ccer_rows])

    grid_factor = float(
        industry_params.get("grid_emission_factor")
        or cost.get("grid_emission_factor")
        or 0.5703
    )
    ccer_ratio = float(compliance_cfg.get("ccer_max_ratio") or 0.05)

    scope1_c = float(emission.scope1_combustion) if emission else 0.0
    scope1_p = float(emission.scope1_process) if emission else 0.0
    scope2 = float(emission.scope2_power) if emission else 0.0
    purchased = float(emission.purchased_mwh) if emission else 0.0

    verified_override = None
    if emission and emission.verified_total is not None:
        verified_override = float(emission.verified_total)
    elif not emission and forecast and forecast.forecast_total:
        verified_override = float(forecast.forecast_total)

    acc = compute_accounting(
        AccountingInput(
            scope1_combustion=scope1_c,
            scope1_process=scope1_p,
            scope2_power=scope2,
            purchased_mwh=purchased,
            market_green_mwh=0.0,
            self_gen_mwh=0.0,
            free_cea_quota=free_cea,
            own_ccer_eligible=own_ccer,
            grid_emission_factor=grid_factor,
            ccer_max_ratio=ccer_ratio,
            verified_override=verified_override,
        )
    )

    cea_market = list_market_cea(db, limit=60)
    prices = [float(r.avg_price) for r in reversed(cea_market)]
    current_cea = prices[-1] if prices else 80.0
    ccer_market = list_market_ccer(db, limit=24)
    current_ccer = float(ccer_market[0].avg_price) if ccer_market else 60.0

    market = judge_market_cycle(
        prices,
        current_cea,
        current_ccer,
        low_percentile=float(compliance_cfg.get("price_low_percentile") or 0.30),
        mid_percentile=float(compliance_cfg.get("price_mid_percentile") or 0.70),
    )

    # 日度碳价预测：用「下一交易日/近期」预测价作为当日无挂单时的锚定
    price_forecast_summary = None
    cea_predicted = None
    ccer_predicted = None
    try:
        from app.services.carbon_compliance.market_sync import fetch_cneeex_daily_quotes_sync
        from app.services.carbon_compliance.price_forecast import forecast_cea_to_year_end

        hist = fetch_cneeex_daily_quotes_sync()
        if not hist:
            hist = [
                {"t": f"{r.year_month}-15", "close": float(r.avg_price)}
                for r in reversed(cea_market)
            ]
        fc = forecast_cea_to_year_end(hist)
        if fc.get("ok"):
            price_forecast_summary = fc.get("summary")
            pts = fc.get("points") or []
            if pts:
                cea_predicted = float(pts[0].get("price") or pts[0].get("close") or current_cea)
            ye = float((price_forecast_summary or {}).get("year_end_price") or current_cea)
            if acc.compliance_gap > 0 and ye > current_cea * 1.03:
                market.rationale = (
                    (market.rationale or "")
                    + f"；日度预测年底约 {ye:.1f} 元/吨（高于现价），宜评估提前分批采购"
                )
            elif acc.compliance_gap < 0 and ye > current_cea * 1.03:
                market.rationale = (
                    (market.rationale or "")
                    + f"；日度预测年底约 {ye:.1f} 元/吨，盈余配额可择高位窗口出售"
                )
    except Exception:
        price_forecast_summary = None

    # CCER 无公开分时：用最新日均作当日预测锚（可后续接独立预测）
    ccer_predicted = current_ccer

    # 结转测算：年末持仓≈清缴后盈余；基础默认取当年免费配额；
    # 净卖出优先取履约年 CEA 台账「当前净卖出」
    net_sell = resolve_cea_net_sell(db, enterprise_id, compliance_year)
    year_end_holding = max(0.0, -min(0.0, float(acc.compliance_gap)))
    carry_base = float(compliance_cfg.get("carry_base_qty") or 0)
    if carry_base <= 0:
        carry_base = allocated_free if allocated_free > 0 else max(0.0, float(acc.free_cea_quota))
    carry = compute_carry_forward(
        base_qty=carry_base,
        net_sell=net_sell,
        year_end_holding=year_end_holding,
        net_sell_multiplier=float(compliance_cfg.get("carry_net_sell_multiplier") or 1.5),
        deadline_md=str(compliance_cfg.get("carry_forward_deadline_md") or "06-10"),
        deadline_year=compliance_year + 1,  # 履约年后的结转日（如 2025 履约 → 2026-06-10）
    )

    payload = build_strategy_payload(
        risk_profile=ent.risk_profile,
        annual_budget_cap=ent.annual_budget_cap,
        single_trade_limit=ent.single_trade_limit,
        settings=settings,
        accounting=acc,
        market=market,
        cea_price=current_cea,
        ccer_price=current_ccer,
        sellable_cea=sellable,
        cea_price_predicted=cea_predicted,
        ccer_price_predicted=ccer_predicted,
        carry_forward=carry,
    )
    if price_forecast_summary:
        payload["price_forecast"] = price_forecast_summary
        tags = dict(payload.get("market_tags") or {})
        tags["price_forecast"] = price_forecast_summary
        payload["market_tags"] = tags
    payload["carry_forward"] = carry.to_dict()
    payload["accounting"] = {
        **payload["accounting"],
        "allocated_free_cea": allocated_free,
        "carry_forward_qty": carry_in,
        "usable_cea_total": float(acc.free_cea_quota),
    }
    tags = dict(payload.get("market_tags") or {})
    tags["carry_forward"] = carry.to_dict()
    payload["market_tags"] = tags
    report_md = strategy_run_to_markdown(
        enterprise_name=ent.name,
        compliance_year=compliance_year,
        accounting=payload["accounting"],
        market_tags=payload["market_tags"],
        plans=payload["plans"],
    )
    run = CarbonStrategyRun(
        enterprise_id=enterprise_id,
        user_id=user_id,
        compliance_year=compliance_year,
        accounting_snapshot=payload["accounting"],
        market_tags=payload["market_tags"],
        plans=payload["plans"],
        status="completed",
        report_md=report_md,
    )
    db.add(run)

    ccer_planned = 0.0
    for p in payload["plans"]:
        for a in p.get("actions") or []:
            if a.get("action") in ("use_ccer", "buy_ccer"):
                ccer_planned = max(ccer_planned, float(a.get("qty") or 0))
                break

    alert_dicts = build_alerts(
        compliance_year=compliance_year,
        clearance_deadline_md=str(industry_params.get("clearance_deadline_md") or "12-31"),
        warn_days=list(compliance_cfg.get("clearance_warn_days") or [90, 30, 15]),
        compliance_gap=acc.compliance_gap,
        ccer_used=ccer_planned,
        ccer_cap=acc.ccer_cap,
        price_band=market.price_band,
        carry_forward=carry,
    )
    for ad in alert_dicts:
        db.add(
            CarbonAlert(
                enterprise_id=enterprise_id,
                user_id=user_id,
                level=ad["level"],
                alert_type=ad["alert_type"],
                message=ad["message"],
                due_at=ad.get("due_at"),
                acked=False,
            )
        )
    db.commit()
    db.refresh(run)

    if notify and alert_dicts:
        try:
            from app.services import notification_service

            top = alert_dicts[0]
            notification_service.create_notification(
                db,
                user_id=user_id,
                title="双碳履约预警",
                body=top["message"],
                link="/system/carbon-assistant",
            )
        except Exception:
            pass

    return run


def list_strategy_runs(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, limit: int = 20
) -> list[CarbonStrategyRun]:
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    return (
        db.query(CarbonStrategyRun)
        .filter(CarbonStrategyRun.enterprise_id == enterprise_id)
        .order_by(CarbonStrategyRun.created_at.desc())
        .limit(limit)
        .all()
    )


def list_alerts(
    db: Session,
    user_id: uuid.UUID,
    *,
    enterprise_id: uuid.UUID | None = None,
    unacked_only: bool = False,
    limit: int = 50,
) -> list[CarbonAlert]:
    q = db.query(CarbonAlert).filter(CarbonAlert.user_id == user_id)
    if enterprise_id:
        q = q.filter(CarbonAlert.enterprise_id == enterprise_id)
    if unacked_only:
        q = q.filter(CarbonAlert.acked.is_(False))
    return q.order_by(CarbonAlert.created_at.desc()).limit(limit).all()


def ack_alert(db: Session, user_id: uuid.UUID, alert_id: uuid.UUID) -> CarbonAlert:
    row = db.get(CarbonAlert, alert_id)
    if not row or row.user_id != user_id:
        raise LookupError("alert not found")
    row.acked = True
    db.commit()
    db.refresh(row)
    return row


def build_import_template_bytes() -> bytes:
    if Workbook is None:
        raise RuntimeError("openpyxl not installed")
    wb = Workbook()
    ws = wb.active
    ws.title = "emissions"
    ws.append(
        [
            "year",
            "scope1_process_万吨",
            "scope1_combustion_万吨",
            "verified_total_万吨",
        ]
    )
    ws.append([2024, 2, 10, None])
    ws2 = wb.create_sheet("cea")
    ws2.append(
        [
            "vintage_year",
            "free_quota_万吨",
            "carry_forward_qty_万吨",
            "net_sell_qty_万吨",
            "avg_cost",
            "sellable_cap_万吨",
        ]
    )
    ws2.append([2024, 14, 0, 0, 0, 1])
    ws3 = wb.create_sheet("ccer")
    ws3.append(
        [
            "project_type",
            "issue_year",
            "expire_at",
            "qty_万吨",
            "cost",
            "eligible_qty_万吨",
        ]
    )
    ws3.append(["wind", 2023, "2030-12-31", 0.5, 50, 0.5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_enterprise_excel(
    db: Session, user_id: uuid.UUID, enterprise_id: uuid.UUID, content: bytes
) -> dict[str, Any]:
    if load_workbook is None:
        raise RuntimeError("openpyxl not installed")
    if not get_enterprise(db, user_id, enterprise_id):
        raise LookupError("enterprise not found")
    wb = load_workbook(io.BytesIO(content), data_only=True)
    counts = {"emissions": 0, "cea": 0, "ccer": 0}
    wan_fields = frozenset(
        {
            "scope1_process",
            "scope1_combustion",
            "verified_total",
            "free_quota",
            "carry_forward_qty",
            "net_sell_qty",
            "sellable_cap",
            "qty",
            "eligible_qty",
        }
    )

    def _normalize_row(raw: dict[str, Any]) -> dict[str, Any]:
        """Excel 数量列按万吨填写；兼容无后缀的旧表头（仍按万吨解释）。"""
        out: dict[str, Any] = {}
        for key, val in raw.items():
            k = str(key or "").strip()
            if k.endswith("_万吨"):
                k = k[: -len("_万吨")]
            if k == "expired_qty":
                k = "carry_forward_qty"
            if k in wan_fields and val is not None and val != "":
                try:
                    out[k] = float(val)
                except (TypeError, ValueError):
                    out[k] = val
            else:
                out[k] = val
        return out

    if "emissions" in wb.sheetnames:
        ws = wb["emissions"]
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            data = _normalize_row(
                {headers[i]: row[i] for i in range(len(headers)) if headers[i]}
            )
            upsert_emission_year(db, user_id, enterprise_id, data)
            counts["emissions"] += 1

    if "cea" in wb.sheetnames:
        ws = wb["cea"]
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            data = _normalize_row(
                {headers[i]: row[i] for i in range(len(headers)) if headers[i]}
            )
            upsert_cea_holding(db, user_id, enterprise_id, data)
            counts["cea"] += 1

    if "ccer" in wb.sheetnames:
        ws = wb["ccer"]
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        for row in rows[1:]:
            if not row or all(c is None for c in row):
                continue
            data = _normalize_row(
                {headers[i]: row[i] for i in range(len(headers)) if headers[i]}
            )
            if data.get("expire_at") is not None:
                data["expire_at"] = str(data["expire_at"])[:10]
            upsert_ccer_holding(db, user_id, enterprise_id, data)
            counts["ccer"] += 1

    return counts


def enterprise_to_dict(e: CarbonEnterprise) -> dict:
    return {
        "id": str(e.id),
        "name": e.name,
        "uscc": e.uscc,
        "industry": e.industry,
        "market_start_year": e.market_start_year,
        "compliance_cycle": e.compliance_cycle,
        "risk_profile": e.risk_profile,
        "annual_budget_cap": e.annual_budget_cap,
        "single_trade_limit": e.single_trade_limit,
        "enterprise_attrs": e.enterprise_attrs or {},
        "created_at": e.created_at.isoformat() if e.created_at else None,
        "updated_at": e.updated_at.isoformat() if e.updated_at else None,
    }


def run_to_dict(r: CarbonStrategyRun) -> dict:
    return {
        "id": str(r.id),
        "enterprise_id": str(r.enterprise_id),
        "compliance_year": r.compliance_year,
        "accounting_snapshot": r.accounting_snapshot,
        "market_tags": r.market_tags,
        "plans": r.plans,
        "status": r.status,
        "report_md": r.report_md,
        "created_at": r.created_at.isoformat() if r.created_at else None,
    }


def alert_to_dict(a: CarbonAlert) -> dict:
    return {
        "id": str(a.id),
        "enterprise_id": str(a.enterprise_id),
        "level": a.level,
        "alert_type": a.alert_type,
        "message": a.message,
        "due_at": a.due_at.isoformat() if a.due_at else None,
        "acked": a.acked,
        "created_at": a.created_at.isoformat() if a.created_at else None,
    }
